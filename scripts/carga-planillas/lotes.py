"""Lotes historicos leidos de las dos planillas de coccion."""
import openpyxl, re, datetime, pathlib

DIR = pathlib.Path("/tmp/claude-0/-home-user/8ac83d0d-fefb-5541-b3aa-b541619709c1/scratchpad/planillas")
GRANDE  = DIR / "56afa5d3-929f243e-8263-4256-9762-1cc11347d077-coccion-grande.xlsx"
PEQUENA = DIR / "c5c61427-d2d9c6f4-9305-4f32-aa44-03f4240eb7af-coccion-pequen-a.xlsx"

# La planilla escribe el mismo estilo de varias formas.
ESTILO = {
    "BLONDE ALE": ("Blonde Ale", "CV-BL-GRANEL"),
    "BLONDE": ("Blonde Ale", "CV-BL-GRANEL"),
    "BONDE ALE": ("Blonde Ale", "CV-BL-GRANEL"),
    "SESSION": ("Session IPA", "CV-SE-GRANEL"),
    "SESSION IPA": ("Session IPA", "CV-SE-GRANEL"),
    "BONDE ALE LIGERA": ("Blonde Ale Ligera", "CV-BLL-GRANEL"),
    "AM. AMBER ALE SUEVE": ("Amber Ale Suave", "CV-AM-GRANEL"),
}

def _fecha(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, (int, float)):      # serial de Excel sin formatear
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
    d, m, a = str(v).split("/")
    return datetime.date(int(a), int(m), int(d))

def leer_grande():
    """Cantidad y Precio son POR UNIDAD DE COMPRA; TOTAL es el de la linea."""
    wb = openpyxl.load_workbook(GRANDE, data_only=True)
    lotes = {}
    for row in wb["Hoja1"].iter_rows(min_row=3, values_only=True):
        if not row[1]: continue
        _, est, lote, fe, ins, cant, precio, total, prov = row[:9]
        k = ("G", int(lote))
        l = lotes.setdefault(k, {"estilo": ESTILO[est.strip().upper()], "fecha": _fecha(fe),
                                 "proveedor": prov, "volumen": 500, "items": [], "costo": 0})
        l["items"].append((ins.strip(), float(cant), float(precio), float(total)))
        l["costo"] += float(total)
    return lotes

def _cant(c):
    s = str(c).strip().replace(",", ".")
    m = re.match(r"^([\d.]+)\s*(kg|g|gr)$", s, re.I)
    if m: return float(m.group(1)), ("kg" if m.group(2).lower() == "kg" else "g")
    return float(s), "un"

def leer_pequena():
    """Cantidad trae la unidad pegada; Precio es el TOTAL de la linea."""
    wb = openpyxl.load_workbook(PEQUENA, data_only=True)
    lotes = {}
    for row in wb["Hoja1"].iter_rows(min_row=3, values_only=True):
        if not row[2]: continue
        fe, lote, est, prov, ins, cant, precio = row[1:8]
        q, u = _cant(cant)
        k = ("P", int(lote))
        l = lotes.setdefault(k, {"estilo": ESTILO[est.strip().upper()], "fecha": _fecha(fe),
                                 "proveedor": prov, "volumen": None, "items": [], "costo": 0})
        l["items"].append((ins.strip(), q, u, float(precio)))
        l["costo"] += float(precio)
    return lotes

def todos():
    out = []
    for (pre, n), l in sorted(leer_grande().items(), key=lambda x: x[0][1]):
        det = " · ".join(f"{i[0]} {i[1]:g} x {i[2]:,.0f}" for i in l["items"])
        out.append(dict(codigo=f"COC-G-{n:02d}", **l, detalle=det))
    # Un lote piloto no produce el mismo granel que uno de 500 L: el estilo es
    # el mismo pero el proceso y el rendimiento no.
    PILOTO = {"CV-BL-GRANEL": "CV-BL-GRANEL-P", "CV-SE-GRANEL": "CV-SE-GRANEL-P"}
    for (pre, n), l in sorted(leer_pequena().items(), key=lambda x: x[0][1]):
        det = " · ".join(f"{i[0]} {i[1]:g}{i[2]} = {i[3]:,.0f}" for i in l["items"])
        nombre, sku = l["estilo"]
        l = dict(l, estilo=(nombre, PILOTO.get(sku, sku)))
        out.append(dict(codigo=f"COC-P-{n:02d}", **l, detalle=det))
    return out

if __name__ == "__main__":
    for l in todos():
        print(f"{l['codigo']}  {l['fecha']}  {l['estilo'][0]:18} {l['costo']:>10,.0f} Gs  "
              f"{len(l['items'])} insumos  vol={l['volumen']}")
